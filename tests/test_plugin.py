import asyncio
import base64
import io
import asyncio
import json
import sys
import tempfile
import types
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import main


class FakeMessage:
    def __init__(self, message):
        self.message = message


class FakeEvent:
    def __init__(self, message, sender="tester", text="", sender_id="u-1", group_id=""):
        self.message_obj = FakeMessage(message)
        self.sender = sender
        self.sender_id = sender_id
        self.group_id = group_id
        self.text = text
        self.sent = []

    def get_sender_id(self):
        return self.sender_id

    def get_sender_name(self):
        return self.sender

    def get_message_id(self):
        return "m-1"

    def get_group_id(self):
        return self.group_id

    def get_message_str(self):
        return self.text

    def plain_result(self, text):
        return text

    async def send(self, message):
        self.sent.append(message)


def make_plugin(root: Path, **config_overrides) -> main.NameSearcherPlugin:
    config = {
        "storage_dir": str(root / "files"),
        "people_file": "",
        "restrict_chat_scope": False,
    }
    config.update(config_overrides)
    return main.NameSearcherPlugin(object(), config)


async def collect_async_generator(generator):
    return [item async for item in generator]


class NameSearcherPluginTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temporary_directory.name)

    def tearDown(self):
        self.temporary_directory.cleanup()

    def test_capture_current_and_forwarded_attachments(self):
        current_file = self.temp_path / "current.txt"
        forwarded_file = self.temp_path / "forwarded.csv"
        current_file.write_text("姓名：张三 班级：软件工程2201", encoding="utf-8")
        forwarded_file.write_text("姓名,专业\n李四,计算机科学", encoding="utf-8")
        event = FakeEvent(
            [
                {"type": "file", "name": current_file.name, "path": str(current_file)},
                {
                    "type": "forward",
                    "nodes": [
                        {
                            "type": "node",
                            "message": [
                                {"type": "file", "name": forwarded_file.name, "path": str(forwarded_file)}
                            ],
                        }
                    ],
                },
            ],
            sender="查找发起人",
        )
        plugin = make_plugin(self.temp_path)

        records = asyncio.run(plugin.capture_event(event))

        self.assertEqual({record["name"] for record in records}, {"current.txt", "forwarded.csv"})
        self.assertEqual(len(plugin.store.list()), 2)
        current_record = next(record for record in records if record["name"] == "current.txt")
        forwarded_record = next(record for record in records if record["name"] == "forwarded.csv")
        self.assertIn("张三", current_record["ocr_text"])
        self.assertIn("李四", forwarded_record["ocr_text"])
        for record in records:
            self.assertEqual(record["source"]["sender_name"], "查找发起人")
            self.assertEqual(record["source"]["message_id"], "m-1")
            self.assertEqual(record["download_status"], "ok")

    def test_zip_capture_imports_all_safe_members_and_nested_archives(self):
        nested_buffer = io.BytesIO()
        with zipfile.ZipFile(nested_buffer, "w", zipfile.ZIP_DEFLATED) as nested:
            nested.writestr("二班名单.txt", "姓名：李四\n班级：软件工程2402".encode("utf-8"))
        archive_path = self.temp_path / "完整名单.zip"
        with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("一班/名单.csv", "姓名,班级\n张三,软件工程2401".encode("utf-8"))
            archive.writestr("嵌套.zip", nested_buffer.getvalue())
            archive.writestr("../越界.txt", "不应导入".encode("utf-8"))
        plugin = make_plugin(
            self.temp_path,
            vision_enabled=False,
            capture_progress_notify=False,
            archive_extraction_enabled=True,
        )

        records = asyncio.run(plugin.capture_event(FakeEvent([
            {"type": "file", "name": archive_path.name, "path": str(archive_path)}
        ])))

        by_name = {record["name"]: record for record in records}
        self.assertIn("完整名单.zip", by_name)
        self.assertIn("名单.csv", by_name)
        self.assertIn("嵌套.zip", by_name)
        self.assertIn("二班名单.txt", by_name)
        self.assertNotIn("越界.txt", by_name)
        self.assertIn("张三", by_name["名单.csv"]["ocr_text"])
        self.assertIn("李四", by_name["二班名单.txt"]["ocr_text"])
        self.assertEqual(by_name["名单.csv"]["source"]["archive_name"], "完整名单.zip")
        self.assertGreaterEqual(by_name["完整名单.zip"]["archive_member_count"], 3)
        self.assertTrue(any("不安全路径" in error for error in by_name["完整名单.zip"]["archive_errors"]))
        self.assertIn("张三", plugin._lookup("张三"))
        self.assertEqual(plugin.latest_progress()["percent"], 100)

    def test_archive_limits_member_count(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("one.txt", "一".encode("utf-8"))
            archive.writestr("two.txt", "二".encode("utf-8"))

        expansion = main._expand_archive_bytes(
            "limited.zip",
            buffer.getvalue(),
            {"archive_extraction_enabled": True, "archive_max_depth": 2, "archive_max_members": 1, "archive_max_total_size_mb": 1},
        )

        self.assertEqual(len(expansion.members), 1)
        self.assertGreaterEqual(expansion.skipped, 1)

    def test_plugin_page_batch_upload_reuses_capture_and_archive_pipeline(self):
        archive_buffer = io.BytesIO()
        with zipfile.ZipFile(archive_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("名单.txt", "姓名：王五\n专业：人工智能".encode("utf-8"))
        plugin = make_plugin(self.temp_path, vision_enabled=False, capture_progress_notify=False)
        payload = {
            "files": [
                {
                    "name": "直接上传.txt",
                    "mime": "text/plain",
                    "data_base64": base64.b64encode("姓名：赵六".encode("utf-8")).decode("ascii"),
                },
                {
                    "name": "上传名单.zip",
                    "mime": "application/zip",
                    "data_base64": base64.b64encode(archive_buffer.getvalue()).decode("ascii"),
                },
            ]
        }

        async def scenario():
            result = await plugin.process_panel_upload(payload)
            tasks = list(plugin._background_tasks)
            if tasks:
                await asyncio.gather(*tasks)
            return result

        result = asyncio.run(scenario())

        self.assertTrue(result["ok"])
        self.assertEqual(result["errors"], [])
        self.assertEqual({record["name"] for record in plugin.store.list()}, {"直接上传.txt", "上传名单.zip", "名单.txt"})
        self.assertIn("赵六", plugin._lookup("赵六"))
        self.assertIn("王五", plugin._lookup("王五"))
        self.assertTrue(all(record["source"]["upload_source"] == "plugin-page" for record in plugin.store.list()))

    def test_panel_upload_returns_queued_before_background_processing_finishes(self):
        plugin = make_plugin(self.temp_path, vision_enabled=False, capture_progress_notify=False)
        payload = {
            "files": [{
                "name": "queued.txt",
                "mime": "text/plain",
                "data_base64": base64.b64encode("姓名：排队测试".encode("utf-8")).decode("ascii"),
            }]
        }

        async def scenario():
            result = await plugin.process_panel_upload(payload)
            self.assertEqual(result["status"], "queued")
            self.assertEqual(result["records"][0]["processing_status"], "queued")
            await asyncio.gather(*list(plugin._background_tasks))
            return result

        result = asyncio.run(scenario())
        self.assertEqual(result["count"], 1)
        self.assertEqual(plugin.store.list()[0]["processing_status"], "completed")
        self.assertIn("排队测试", plugin._lookup("排队测试"))

    def test_pending_panel_upload_resumes_after_plugin_restart(self):
        plugin = make_plugin(self.temp_path, vision_enabled=False, capture_progress_notify=False)
        record = plugin.store.add_bytes(
            "restart.txt",
            "姓名：重启恢复".encode("utf-8"),
            kind="document",
            mime="text/plain",
            source={"upload_source": "plugin-page"},
        )
        plugin.store.update(
            record["id"],
            processing_status="queued",
            processing_updated_at=main._now(),
            download_status="ok",
        )
        reloaded = make_plugin(self.temp_path, vision_enabled=False, capture_progress_notify=False)

        async def scenario():
            reloaded._schedule_pending_uploads()
            await asyncio.gather(*list(reloaded._background_tasks))

        asyncio.run(scenario())
        restored = reloaded.store.get(record["id"])
        self.assertEqual(restored["processing_status"], "completed")
        self.assertIn("重启恢复", reloaded._lookup("重启恢复"))

    def test_archive_resume_skips_completed_member_and_continues_next_file(self):
        plugin = make_plugin(
            self.temp_path,
            vision_enabled=False,
            llm_review_enabled=False,
            capture_progress_notify=False,
            archive_extraction_enabled=True,
        )
        archive_buffer = io.BytesIO()
        with zipfile.ZipFile(archive_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("first.txt", "姓名：已完成".encode("utf-8"))
            archive.writestr("second.txt", "姓名：待恢复".encode("utf-8"))
        source = {"upload_source": "plugin-page"}
        parent = plugin.store.add_bytes(
            "resume.zip",
            archive_buffer.getvalue(),
            kind="archive",
            mime="application/zip",
            source=source,
        )
        plugin.store.update(parent["id"], processing_status="queued")
        first = plugin.store.add_bytes(
            "resume.zip/first.txt",
            "姓名：已完成".encode("utf-8"),
            kind="document",
            mime="text/plain",
            source={**source, "archive_id": parent["id"], "archive_member": "first.txt"},
        )
        plugin.store.update(
            first["id"],
            processing_status="completed",
            extraction_status="extracted",
            ocr_text="姓名：已完成",
        )
        second = plugin.store.add_bytes(
            "resume.zip/second.txt",
            "姓名：待恢复".encode("utf-8"),
            kind="document",
            mime="text/plain",
            source={**source, "archive_id": parent["id"], "archive_member": "second.txt"},
        )
        plugin.store.update(second["id"], processing_status="queued")

        reloaded = make_plugin(
            self.temp_path,
            vision_enabled=False,
            llm_review_enabled=False,
            capture_progress_notify=False,
            archive_extraction_enabled=True,
        )
        captured_names = []
        original_capture_one = reloaded._capture_one

        def tracked_capture(candidate):
            captured_names.append(candidate.name)
            return original_capture_one(candidate)

        reloaded._capture_one = tracked_capture

        async def scenario():
            reloaded._schedule_pending_uploads()
            await asyncio.gather(*list(reloaded._background_tasks))

        asyncio.run(scenario())
        self.assertNotIn("first.txt", captured_names)
        self.assertEqual(captured_names.count("second.txt"), 1)
        self.assertEqual(reloaded.store.get(first["id"])["processing_status"], "completed")
        self.assertEqual(reloaded.store.get(second["id"])["processing_status"], "completed")
        self.assertIn("待恢复", reloaded._lookup("待恢复"))

    def test_public_file_record_excludes_large_model_evidence(self):
        plugin = make_plugin(self.temp_path)
        record = plugin.store.add_bytes("large.txt", b"content", kind="document", mime="text/plain", source={})
        plugin.store.update(
            record["id"],
            vision_results=[{"text": "x" * 500_000}],
            vision_consensus_text="y" * 500_000,
            llm_review_text="z" * 500_000,
            ocr_text_local="w" * 500_000,
        )
        public = main._public_record(plugin.store.get(record["id"]))
        self.assertNotIn("vision_results", public)
        self.assertNotIn("vision_consensus_text", public)
        self.assertNotIn("llm_review_text", public)
        self.assertNotIn("ocr_text_local", public)
        self.assertLess(len(json.dumps(public, ensure_ascii=False)), 10_000)

    def test_capture_deduplicates_same_content_across_current_and_forwarded_files(self):
        first_file = self.temp_path / "first.txt"
        second_file = self.temp_path / "second.txt"
        content = "姓名：王五 专业：人工智能"
        first_file.write_text(content, encoding="utf-8")
        second_file.write_text(content, encoding="utf-8")
        event = FakeEvent(
            [
                {"type": "file", "name": first_file.name, "path": str(first_file)},
                {
                    "type": "forward",
                    "content": [{"type": "file", "name": second_file.name, "path": str(second_file)}],
                },
            ]
        )
        plugin = make_plugin(self.temp_path)

        records = asyncio.run(plugin.capture_event(event))

        self.assertEqual(len(records), 1)
        self.assertEqual(len(plugin.store.list()), 1)
        persisted = main.ArtifactStore(self.temp_path / "files").list()
        self.assertEqual(len(persisted), 1)
        self.assertEqual(persisted[0]["sha256"], records[0]["sha256"])
        progress = plugin.latest_progress()
        self.assertEqual(progress["percent"], 100)
        self.assertEqual(progress["completed"], 2)
        self.assertEqual(len(progress["records"]), 1)

    def test_multiselect_forward_and_qq_image_segment_are_captured(self):
        first = self.temp_path / "qq-image-one.jpg"
        second = self.temp_path / "qq-image-two.jpg"
        first.write_bytes(b"first-image")
        second.write_bytes(b"second-image")
        event = FakeEvent([])
        event.selected_messages = [
            {"type": "node", "content": [{"type": "image", "name": first.name, "file": "cache-token", "url": str(first)}]},
            {"type": "node", "content": [{"type": "image", "name": second.name, "file": "cache-token", "url": str(second)}]},
        ]
        plugin = make_plugin(self.temp_path, capture_progress_notify=False)
        recognition = {"text": "名单", "confidence": 0.9, "clarity": 0.9, "engine": "test", "data": {}}

        with patch.object(main, "_recognize_image", return_value=recognition):
            records = asyncio.run(plugin.capture_event(event))

        self.assertEqual({item["name"] for item in records}, {first.name, second.name})
        self.assertTrue(all(item["kind"] == "image" for item in records))
        progress = plugin.latest_progress()
        self.assertEqual(progress["percent"], 100)
        self.assertEqual({item["name"] for item in progress["records"]}, {first.name, second.name})
        output = asyncio.run(collect_async_generator(plugin.capture_progress(FakeEvent([], text="/捕获进度"))))
        self.assertIn("100%", output[0])
        self.assertIn(first.name, output[0])

    def test_qq_forward_id_payload_is_fetched_before_capture(self):
        image = self.temp_path / "forwarded-qq-image.jpg"
        image.write_bytes(b"forward-image")
        event = FakeEvent([{"type": "forward", "id": "forward-123"}])

        async def get_forward_msg(message_id):
            self.assertEqual(message_id, "forward-123")
            return {"messages": [{"content": [{"type": "image", "name": image.name, "url": str(image)}]}]}

        event.get_forward_msg = get_forward_msg
        plugin = make_plugin(self.temp_path, capture_progress_notify=False)
        recognition = {"text": "张三", "confidence": 0.9, "clarity": 0.9, "engine": "test", "data": {}}

        with patch.object(main, "_recognize_image", return_value=recognition):
            records = asyncio.run(plugin.capture_event(event))

        self.assertEqual([item["name"] for item in records], [image.name])
        self.assertEqual(plugin.latest_progress()["percent"], 100)

    def test_qq_napcat_forward_file_id_is_resolved_through_bot(self):
        content = self.temp_path / "forwarded-roster.docx"
        content.write_bytes(b"forwarded-document")

        class Bot:
            def __init__(self):
                self.calls = []

            async def call_action(self, *, action, **kwargs):
                self.calls.append((action, kwargs))
                if action == "get_forward_msg":
                    return {"data": {"messages": [{"sender": {"user_id": 42, "nickname": "原发送者"}, "message": [{"type": "file", "data": {"file_id": "file-1", "file": content.name}}]}]}}
                if action == "get_group_file_url":
                    return {"url": str(content), "file_name": content.name}
                raise AssertionError(action)

        event = FakeEvent([{"type": "forward", "data": {"id": "forward-1"}}], group_id="10001")
        event.bot = Bot()
        plugin = make_plugin(self.temp_path, capture_progress_notify=False)

        records = asyncio.run(plugin.capture_event(event))

        self.assertEqual([item["name"] for item in records], [content.name])
        self.assertEqual(records[0]["source"]["sender_name"], "原发送者")
        self.assertEqual([item[0] for item in event.bot.calls], ["get_forward_msg", "get_group_file_url"])

    def test_clear_failed_only_removes_unavailable_records(self):
        plugin = make_plugin(self.temp_path)
        kept = plugin.store.add_bytes("kept.txt", b"ok", kind="document", mime="text/plain", source={})
        failed = plugin.store.add_unavailable("failed.pdf", kind="document", mime="application/pdf", source={}, reason="download_failed")

        deleted = plugin.store.clear_failed()

        self.assertEqual(deleted, [failed["id"]])
        self.assertIsNotNone(plugin.store.get(kept["id"]))
        self.assertIsNone(plugin.store.get(failed["id"]))

    def test_restrict_chat_scope_is_enabled_by_default(self):
        plugin = main.NameSearcherPlugin(object(), {"storage_dir": str(self.temp_path / "files")})
        self.assertTrue(plugin.config["restrict_chat_scope"])
        self.assertFalse(plugin.is_event_allowed(FakeEvent([], sender_id="not-configured")))

    def test_admin_only_capture_filters_by_event_sender_without_blocking_commands(self):
        source = self.temp_path / "admin-only.txt"
        source.write_text("姓名：张三", encoding="utf-8")
        plugin = make_plugin(
            self.temp_path,
            admin_only_capture=True,
            admin_ids="admin-1，admin-2; admin-3",
            capture_progress_notify=False,
        )
        attachment = [{"type": "file", "name": source.name, "path": str(source)}]

        allowed = asyncio.run(plugin.capture_event(FakeEvent(attachment, sender_id="admin-2")))
        denied = asyncio.run(plugin.capture_event(FakeEvent(attachment, sender_id="user-1")))
        command = asyncio.run(collect_async_generator(plugin.find(FakeEvent([], text="/查找 张三", sender_id="user-1"))))

        self.assertEqual(len(allowed), 1)
        self.assertEqual(denied, [])
        self.assertEqual(len(command), 1)

    def test_admin_only_capture_with_empty_admin_list_captures_nothing(self):
        source = self.temp_path / "blocked.txt"
        source.write_text("姓名：李四", encoding="utf-8")
        plugin = make_plugin(self.temp_path, admin_only_capture=True, admin_ids="")

        records = asyncio.run(plugin.capture_event(FakeEvent(
            [{"type": "file", "name": source.name, "path": str(source)}],
            sender_id="admin-1",
        )))

        self.assertEqual(records, [])
        self.assertEqual(plugin.store.list(), [])

    def test_vision_provider_can_replace_local_ocr_text(self):
        class VisionResponse:
            completion_text = "姓名 张三 班级 软件工程2201"

        class VisionProvider:
            async def text_chat(self, prompt, session_id, image_urls, **kwargs):
                self.image_urls = image_urls
                return VisionResponse()

        class VisionContext:
            def __init__(self):
                self.provider = VisionProvider()

            def get_using_provider(self):
                return self.provider

        image = self.temp_path / "vision.jpg"
        image.write_bytes(b"test-image")
        context = VisionContext()
        plugin = main.NameSearcherPlugin(context, {
            "storage_dir": str(self.temp_path / "files"),
            "restrict_chat_scope": False,
            "vision_enabled": True,
            "capture_progress_notify": False,
        })
        local = {"text": "local", "confidence": 0.3, "clarity": 0.8, "engine": "test", "data": {}}

        with patch.object(main, "_recognize_image", return_value=local):
            records = asyncio.run(plugin.capture_event(FakeEvent([{"type": "image", "name": image.name, "path": str(image)}])))

        self.assertIn("张三", records[0]["ocr_text"])
        self.assertEqual(records[0]["ocr_engine"], "vision-model")
        self.assertTrue(context.provider.image_urls[0].startswith("data:image/jpeg;base64,"))

    def test_current_astrbot_llm_generate_api_is_preferred_for_vision(self):
        class VisionResponse:
            completion_text = "姓名 李四 专业 人工智能"

        class VisionContext:
            def __init__(self):
                self.calls = []

            async def get_current_chat_provider_id(self, umo):
                self.umo = umo
                return "vision-provider-id"

            async def llm_generate(self, **kwargs):
                self.calls.append(kwargs)
                return VisionResponse()

        image = self.temp_path / "official-vision.jpg"
        image.write_bytes(b"test-image")
        context = VisionContext()
        plugin = main.NameSearcherPlugin(context, {
            "storage_dir": str(self.temp_path / "files"),
            "restrict_chat_scope": False,
            "vision_enabled": True,
            "capture_progress_notify": False,
        })
        local = {"text": "local", "confidence": 0.3, "clarity": 0.8, "engine": "test", "data": {}}
        event = FakeEvent([{"type": "image", "name": image.name, "path": str(image)}])
        event.unified_msg_origin = "aiocqhttp:group:123"

        with patch.object(main, "_recognize_image", return_value=local):
            records = asyncio.run(plugin.capture_event(event))

        self.assertIn("李四", records[0]["ocr_text"])
        self.assertEqual(context.umo, event.unified_msg_origin)
        self.assertEqual(context.calls[0]["chat_provider_id"], "vision-provider-id")
        self.assertTrue(context.calls[0]["image_urls"][0].startswith("data:image/jpeg;base64,"))

    def test_lookup_from_capture_includes_file_resource(self):
        plugin = make_plugin(self.temp_path)
        record = plugin.store.add_bytes("名单.xlsx", b"content", kind="spreadsheet", mime="application/octet-stream", source={})
        plugin.store.update(record["id"], ocr_text="张三 | 软件工程2201")

        output = plugin._lookup("张三")

        self.assertIn("文件资源：名单.xlsx", output)
        self.assertIn(f"/api/name-searcher/files/{record['id']}/content", output)

    def test_lookup_searches_every_retained_vision_result(self):
        plugin = make_plugin(self.temp_path)
        record = plugin.store.add_bytes("视觉名单.png", b"image", kind="image", mime="image/png", source={})
        plugin.store.update(
            record["id"],
            ocr_text="视觉共识仅保留：李四",
            vision_consensus_text="姓名：李四",
            vision_results=[
                {
                    "provider": "provider-a",
                    "model": "vision-a",
                    "label": "provider-a::vision-a",
                    "status": "completed",
                    "text": "1. **张\u200b三** | 班级：软件工程2401\n2. 李四",
                },
                {
                    "provider": "provider-b",
                    "model": "vision-b",
                    "status": "completed",
                    "text": "只识别到李四",
                },
            ],
        )

        reloaded = make_plugin(self.temp_path)
        matches = reloaded._lookup_matches("张三")
        output = reloaded._lookup("张三")

        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0]["source"]["id"], record["id"])
        self.assertEqual(matches[0]["record"]["匹配依据"], "视觉模型 provider-a::vision-a")
        self.assertIn("张三", output)
        self.assertIn("视觉模型 provider-a::vision-a", output)

    def test_lookup_normalises_escaped_and_punctuated_model_names(self):
        plugin = make_plugin(self.temp_path)
        record = plugin.store.add_bytes("escaped.txt", b"content", kind="document", mime="text/plain", source={})
        plugin.store.update(
            record["id"],
            ocr_text="当前最终结果没有目标姓名",
            vision_results=[{
                "label": "escaped-model",
                "text": r"姓名：\u738b·\u4e94；专业：人工智能",
            }],
        )

        matches = plugin._lookup_matches("王五")

        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0]["record"]["匹配依据"], "视觉模型 escaped-model")

    def test_lookup_includes_people_details_and_source(self):
        people_file = self.temp_path / "people.json"
        people_file.write_text(
            json.dumps(
                {
                    "people": [
                        {
                            "姓名": "张三",
                            "班级": "软件工程2201",
                            "专业": "软件工程",
                            "学号": "20220001",
                        }
                    ]
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        plugin = make_plugin(self.temp_path, people_file=str(people_file))

        output = plugin._lookup("张三")

        self.assertIn("查找对象：张三", output)
        self.assertIn("班级：软件工程2201", output)
        self.assertIn("专业：软件工程", output)
        self.assertIn("学号：20220001", output)
        self.assertIn(f"来源：{people_file}", output)
        self.assertNotIn("**", output)

    def _asgi(self, app, method, path, payload=None, headers=None, query=b""):
        body = json.dumps(payload).encode("utf-8") if payload is not None else b""
        sent = []
        delivered = False

        async def receive():
            nonlocal delivered
            if delivered:
                return {"type": "http.disconnect"}
            delivered = True
            return {"type": "http.request", "body": body, "more_body": False}

        async def send(message):
            sent.append(message)

        scope = {
            "type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1",
            "method": method, "scheme": "http", "path": path, "raw_path": path.encode("utf-8"),
            "query_string": query, "root_path": "",
            "headers": [(b"content-type", b"application/json")] + list(headers or []),
            "client": ("127.0.0.1", 1), "server": ("test", 80),
        }
        asyncio.run(app(scope, receive, send))
        status = next(m["status"] for m in sent if m["type"] == "http.response.start")
        raw = b"".join(m.get("body", b"") for m in sent if m["type"] == "http.response.body")
        try:
            return status, json.loads(raw or b"{}")
        except ValueError:
            return status, raw

    def test_legacy_panel_refuses_api_without_configured_token(self):
        plugin = make_plugin(self.temp_path)
        plugin.store.add_bytes("名单.txt", "张三".encode("utf-8"), kind="document", mime="text/plain", source={})
        app = main._build_web_app(plugin)
        for method, path in (("GET", "/files"), ("GET", "/config"), ("POST", "/files/clear"), ("PATCH", "/config"), ("GET", "/api/name-searcher/files")):
            status, _ = self._asgi(app, method, path, {} if method != "GET" else None)
            self.assertEqual(status, 403, path)
        self.assertEqual(len(plugin.store.list()), 1)

    def test_legacy_panel_rejects_wrong_token_and_accepts_right_one(self):
        plugin = make_plugin(self.temp_path, panel_access_token="right")
        app = main._build_web_app(plugin)
        status, _ = self._asgi(app, "GET", "/files", headers=[(b"authorization", b"Bearer wrong")])
        self.assertEqual(status, 401)
        status, _ = self._asgi(app, "GET", "/files", query=b"token=right")
        self.assertEqual(status, 200)
        status, _ = self._asgi(app, "GET", "/files", headers=[(b"x-namesearcher-token", b"right")])
        self.assertEqual(status, 200)

    def test_panel_config_cannot_change_locked_keys_or_read_token(self):
        secret = self.temp_path / "secret.json"
        secret.write_text(json.dumps([{"姓名": "机密", "电话": "123"}], ensure_ascii=False), encoding="utf-8")
        plugin = make_plugin(self.temp_path, panel_access_token="tok")
        original_storage = plugin.config["storage_dir"]
        ignored = main._apply_panel_config(plugin, {
            "people_file": str(secret),
            "storage_dir": "/tmp/elsewhere",
            "panel_access_token": "",
            "merged_forward_results": False,
        })
        self.assertEqual(set(ignored), {"people_file", "storage_dir", "panel_access_token"})
        self.assertEqual(plugin.config["people_file"], "")
        self.assertEqual(plugin.config["storage_dir"], original_storage)
        self.assertEqual(plugin.config["panel_access_token"], "tok")
        self.assertFalse(plugin.config["merged_forward_results"])
        self.assertIn("未找到", plugin._lookup("机密"))
        view = main._panel_config_view(plugin)
        self.assertNotIn("panel_access_token", view)
        self.assertTrue(view["panel_access_token_set"])

    def test_people_lookup_does_not_match_different_similar_names(self):
        people_file = self.temp_path / "people.json"
        people_file.write_text(json.dumps([
            {"姓名": "王小红", "班级": "A"},
            {"姓名": "李明华", "班级": "B"},
            {"姓名": "张四", "班级": "C"},
        ], ensure_ascii=False), encoding="utf-8")
        plugin = make_plugin(self.temp_path, people_file=str(people_file))
        self.assertEqual(plugin._lookup_matches("王小明"), [])
        self.assertEqual(plugin._lookup_matches("李明"), [])
        self.assertEqual(plugin._lookup_matches("张三"), [])
        self.assertEqual(len(plugin._lookup_matches("王小红")), 1)

    def test_file_lookup_prefers_whole_name_over_longer_name(self):
        plugin = make_plugin(self.temp_path)
        longer = plugin.store.add_bytes("a.txt", b"a", kind="document", mime="text/plain", source={})
        plugin.store.update(longer["id"], ocr_text="第 1 行 | 张三丰 | 武当派")
        whole = plugin.store.add_bytes("b.txt", b"b", kind="document", mime="text/plain", source={})
        plugin.store.update(whole["id"], ocr_text="第 2 行 | 张三 | 软件工程")
        matches = plugin._lookup_matches("张三")
        self.assertEqual([m["source"]["id"] for m in matches], [whole["id"]])
        output = plugin._lookup("张三")
        self.assertIn("另有 1 条", output)

    def test_file_lookup_falls_back_to_partial_match_with_warning(self):
        plugin = make_plugin(self.temp_path)
        record = plugin.store.add_bytes("a.txt", b"a", kind="document", mime="text/plain", source={})
        plugin.store.update(record["id"], ocr_text="张三丰 武当派")
        matches = plugin._lookup_matches("张三")
        self.assertEqual(len(matches), 1)
        self.assertEqual(matches[0]["strength"], "partial")
        self.assertIn("没有完全一致的姓名", plugin._lookup("张三"))

    def test_file_lookup_accepts_labelled_and_punctuated_names(self):
        plugin = make_plugin(self.temp_path)
        record = plugin.store.add_bytes("a.txt", b"a", kind="document", mime="text/plain", source={})
        plugin.store.update(record["id"], ocr_text="姓名张三，班级软件2201\n赵 六 数据科学")
        self.assertEqual(plugin._lookup_matches("张三")[0]["strength"], "exact")
        self.assertEqual(plugin._lookup_matches("赵六")[0]["strength"], "exact")

    def test_capture_window_turns_off_at_zero_seconds(self):
        plugin = make_plugin(self.temp_path, capture_interval_seconds=0)
        self.assertFalse(plugin._capture_window_on())
        plugin.config["capture_interval_seconds"] = 15
        self.assertTrue(plugin._capture_window_on())
        self.assertEqual(plugin._capture_interval(), 15)
        plugin.config["capture_window_enabled"] = False
        self.assertFalse(plugin._capture_window_on())

    def test_duplicate_knobs_hidden_from_astrbot_config_editor(self):
        schema = json.loads((Path(main.__file__).parent / "_conf_schema.json").read_text(encoding="utf-8"))
        for key in ("vision_provider", "vision_model", "capture_window_enabled", "clear_downloaded_files"):
            self.assertTrue(schema[key].get("invisible"), key)
        self.assertTrue(schema["panel_access_token"].get("secret"))
        page = (Path(main.__file__).parent / "pages" / "name-searcher" / "index.html").read_text(encoding="utf-8")
        for removed in ("retentionInput", "configClearButton", 'id="visionProvider"', 'id="llmReviewProvider"', 'id="captureWindowEnabled"'):
            self.assertNotIn(removed, page)
        # v1.5.x split model targets on the letters r/n because of a doubled backslash.
        self.assertNotIn("[\\\\r\\\\n", page)

    def test_plugin_page_has_type_source_filters_and_collapse(self):
        root = Path(main.__file__).parent
        pages = [(root / rel).read_text(encoding="utf-8") for rel in ("pages/name-searcher/index.html", "webui/index.html", "plugin-page/index.html")]
        self.assertEqual(len(set(pages)), 1)
        for marker in ('id="typeFilter"', 'id="sourceFilter"', 'value="latest"', "COLLAPSE_AT = 5", "data-collapse"):
            self.assertIn(marker, pages[0])

    def _make_bundle(self):
        import subprocess, sys
        folder = self.temp_path / "in"
        folder.mkdir()
        (folder / "名单.csv").write_text("姓名,班级\n王小明,软件2201\n", encoding="utf-8")
        (folder / "备注.txt").write_text("张三 计算机2202", encoding="utf-8")
        tool = Path(main.__file__).parent / "tools" / "name_searcher_local.py"
        out = self.temp_path / "bundle.zip"
        subprocess.run([sys.executable, str(tool), "run", str(folder), "--work", str(self.temp_path / "work"), "--out", str(out)], check=True, capture_output=True)
        return out

    def test_local_tool_bundle_imports_in_chunks_and_is_searchable(self):
        import base64
        bundle = self._make_bundle().read_bytes()
        plugin = make_plugin(self.temp_path)
        offset = 0
        for start in range(0, len(bundle), 1000):
            result, status = plugin.import_chunk({"upload_id": "abcdef123456", "offset": offset, "data_base64": base64.b64encode(bundle[start:start + 1000]).decode()})
            self.assertEqual(status, 200, result)
            offset = result["received"]
        summary, status = plugin.import_staged("abcdef123456", "测试包")
        self.assertEqual(status, 200)
        self.assertEqual(summary["created"], 2)
        self.assertTrue(all(item.get("processing_status") == "completed" for item in plugin.store.list()))
        self.assertTrue(all((item.get("source") or {}).get("upload_source") == "local-tool" for item in plugin.store.list()))
        self.assertIn("王小明", plugin._lookup("王小明"))
        again, _ = plugin.import_staged("abcdef123456")
        self.assertFalse(again.get("ok", True))
        result, status = plugin.import_chunk({"upload_id": "abcdef654321", "offset": 0, "data_base64": base64.b64encode(bundle).decode()})
        summary, _ = plugin.import_staged("abcdef654321")
        self.assertEqual((summary["created"], summary["exists"]), (0, 2))

    def test_bundle_import_rejects_bad_ids_order_and_tampering(self):
        import zipfile
        plugin = make_plugin(self.temp_path)
        self.assertEqual(plugin.import_chunk({"upload_id": "../../etc", "offset": 0, "data_base64": ""})[1], 400)
        self.assertEqual(plugin.import_chunk({"upload_id": "abcdef123456", "offset": 9, "data_base64": "AA=="})[1], 409)
        source = self._make_bundle()
        tampered = plugin.import_inbox() / "tampered.zip"
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(tampered, "w") as copy:
            for name in original.namelist():
                data = original.read(name)
                copy.writestr(name, data + b"!" if name.startswith("files/") else data)
        result = plugin.import_from_inbox()
        self.assertEqual(result["bundles"][0]["created"], 0)
        self.assertEqual(len(result["bundles"][0]["errors"]), 2)
        self.assertTrue((plugin.import_inbox() / "done" / "tampered.zip").exists())
        not_bundle = plugin.import_inbox() / "x.zip"
        not_bundle.write_bytes(b"not a zip")
        self.assertFalse(plugin.import_from_inbox()["bundles"][0]["ok"])

    def test_upload_interrupted_twice_is_skipped(self):
        plugin = make_plugin(self.temp_path)
        record = plugin.store.add_bytes("x.txt", b"hello", kind="document", mime="text/plain", source={"upload_source": "plugin-page"})
        plugin.store.update(record["id"], processing_status="running", processing_attempts=2, extraction_status="pending")
        asyncio.run(plugin._process_stored_upload(record["id"]))
        updated = plugin.store.get(record["id"])
        self.assertEqual(updated["processing_status"], "failed")
        self.assertIn("多次被中断", updated["processing_error"])

    def test_batch_find_command_is_removed(self):
        self.assertFalse(hasattr(main.NameSearcherPlugin, "batch_find"))

    def test_panel_config_can_toggle_merged_forward_results(self):
        plugin = make_plugin(self.temp_path)

        main._apply_panel_config(plugin, {"merged_forward_results": True})

        self.assertTrue(plugin.config["merged_forward_results"])

    def test_panel_config_accepts_vision_targets_and_capture_window(self):
        plugin = make_plugin(self.temp_path)

        main._apply_panel_config(plugin, {
            "vision_model_targets": "provider-a::vision-a\nprovider-b::vision-b",
            "vision_review_max_pages": 6,
            "vision_consensus_enabled": True,
            "capture_window_enabled": True,
            "capture_interval_seconds": 45,
        })

        self.assertEqual(main._parse_vision_targets(plugin.config), [
            ("provider-a", "vision-a"),
            ("provider-b", "vision-b"),
        ])
        self.assertEqual(plugin.config["vision_review_max_pages"], 6)
        self.assertEqual(plugin.config["capture_interval_seconds"], 45)

    def test_merged_forward_results_are_enabled_by_default(self):
        plugin = make_plugin(self.temp_path)
        self.assertTrue(plugin.config["merged_forward_results"])

    def test_find_can_return_all_matches_as_merged_forward_nodes(self):
        class Plain:
            def __init__(self, text):
                self.text = text

        class Node:
            def __init__(self, content, **kwargs):
                self.content = content
                self.name = kwargs.get("name")

        class Nodes:
            def __init__(self, nodes):
                self.nodes = nodes

        class File:
            def __init__(self, name, file="", url=""):
                self.name = name
                self.file = file

        class Image:
            @classmethod
            def fromFileSystem(cls, path):
                item = cls()
                item.path = path
                return item

        components = types.ModuleType("astrbot.api.message_components")
        components.Plain = Plain
        components.Node = Node
        components.Nodes = Nodes
        components.File = File
        components.Image = Image
        astrbot = types.ModuleType("astrbot")
        api = types.ModuleType("astrbot.api")
        astrbot.api = api
        api.message_components = components

        plugin = make_plugin(self.temp_path, merged_forward_results=True)
        matches = [
            {"record": {"姓名": "张三", "班级": f"软件工程{i}"}, "score": 1.0, "source": {"name": "人员库"}}
            for i in range(10)
        ]
        plugin._lookup_matches = lambda _query: matches
        event = FakeEvent([], text="/查找 张三")
        event.chain_result = lambda chain: chain

        with patch.dict(sys.modules, {
            "astrbot": astrbot,
            "astrbot.api": api,
            "astrbot.api.message_components": components,
        }):
            output = asyncio.run(collect_async_generator(plugin.find(event)))

        self.assertEqual(len(output), 1)
        merged = output[0][0]
        self.assertIsInstance(merged, Nodes)
        self.assertEqual(len(merged.nodes), 11)
        self.assertIn("匹配数量：10", merged.nodes[0].content[0].text)
        self.assertIn("结果 10", merged.nodes[-1].content[0].text)

    def test_merged_forward_results_fall_back_on_unsupported_platform(self):
        plugin = make_plugin(self.temp_path, merged_forward_results=True)
        event = FakeEvent([], text="/查找 张三")
        event.chain_result = lambda chain: chain
        event.get_platform_name = lambda: "telegram"

        output = asyncio.run(collect_async_generator(plugin.find(event)))

        self.assertEqual(len(output), 1)
        self.assertIsInstance(output[0], str)

    def test_filter_uses_event_message_type_from_filter_namespace(self):
        self.assertIs(main.EventMessageType, main.filter.EventMessageType)

    def test_ocr_data_preserves_lines_and_confidence(self):
        data = {
            "text": ["张三", "软件工程", "李四"],
            "conf": ["90", "80", "100"],
            "left": [10, 80, 10],
            "page_num": [1, 1, 1],
            "block_num": [1, 1, 1],
            "par_num": [1, 1, 1],
            "line_num": [1, 1, 2],
        }

        text, confidence, word_count = main._ocr_text_from_data(data)

        self.assertEqual(text, "张三 软件工程\n李四")
        self.assertAlmostEqual(confidence, 0.9)
        self.assertEqual(word_count, 3)

    def test_rapidocr_results_feed_common_ocr_pipeline(self):
        rapid_result = [
            [[[80, 30], [150, 30], [150, 50], [80, 50]], "软件工程", 0.86],
            [[[10, 10], [60, 10], [60, 28], [10, 28]], "张三", 0.94],
        ]

        data = main._rapid_result_to_data(rapid_result)
        text, confidence, word_count = main._ocr_text_from_data(data)

        self.assertEqual(text, "张三\n软件工程")
        self.assertAlmostEqual(confidence, 0.9)
        self.assertEqual(word_count, 2)

    def test_rapidocr_groups_cells_on_the_same_visual_row(self):
        rapid_result = [
            [[[10, 10], [60, 10], [60, 30], [10, 30]], "张三", 0.94],
            [[[100, 12], [180, 12], [180, 32], [100, 32]], "软件工程", 0.86],
            [[[10, 48], [60, 48], [60, 68], [10, 68]], "李四", 0.92],
            [[[100, 50], [180, 50], [180, 70], [100, 70]], "人工智能", 0.88],
        ]

        data = main._rapid_result_to_data(rapid_result)
        text, confidence, word_count = main._ocr_text_from_data(data)

        self.assertEqual(text, "张三 软件工程\n李四 人工智能")
        self.assertAlmostEqual(confidence, 0.9)
        self.assertEqual(word_count, 4)

    def test_csv_and_tsv_are_classified_as_spreadsheets(self):
        self.assertEqual(main._kind_for("students.csv", "text/csv"), "spreadsheet")
        self.assertEqual(main._kind_for("students.tsv", "text/tab-separated-values"), "spreadsheet")
        self.assertEqual(main._kind_for("portrait.jfif", "image/jpeg"), "image")

    def test_chat_scope_limits_capture_and_commands(self):
        source = self.temp_path / "allowed.txt"
        source.write_text("姓名：张三", encoding="utf-8")
        plugin = make_plugin(
            self.temp_path,
            restrict_chat_scope=True,
            allowed_group_ids="group-1, group-2",
            allowed_private_ids="user-1；user-2",
        )

        allowed_group = FakeEvent(
            [{"type": "file", "name": source.name, "path": str(source)}],
            group_id="group-2",
            sender_id="someone",
        )
        denied_group = FakeEvent(
            [{"type": "file", "name": source.name, "path": str(source)}],
            group_id="group-3",
            sender_id="someone",
        )
        self.assertTrue(plugin.is_event_allowed(allowed_group))
        self.assertFalse(plugin.is_event_allowed(denied_group))
        self.assertEqual(len(asyncio.run(plugin.capture_event(allowed_group))), 1)
        self.assertEqual(asyncio.run(plugin.capture_event(denied_group)), [])

        allowed_private = FakeEvent([], text="/查找 张三", sender_id="user-1")
        denied_private = FakeEvent([], text="/查找 张三", sender_id="user-9")
        self.assertTrue(plugin.is_event_allowed(allowed_private))
        self.assertFalse(plugin.is_event_allowed(denied_private))
        self.assertEqual(len(asyncio.run(collect_async_generator(plugin.find(allowed_private)))), 1)
        self.assertEqual(asyncio.run(collect_async_generator(plugin.find(denied_private))), [])

    def test_low_confidence_image_requests_llm_review_or_clearer_copy(self):
        image_file = self.temp_path / "blurred.png"
        image_file.write_bytes(b"not-a-real-image")
        event = FakeEvent(
            [{"type": "image", "name": image_file.name, "path": str(image_file)}],
            sender="审核发起人",
        )
        plugin = make_plugin(self.temp_path, review_threshold=0.55)

        recognition = {
            "text": "模糊文字",
            "confidence": 0.2,
            "raw_confidence": 0.3,
            "clarity": 0.15,
            "width": 640,
            "height": 480,
            "engine": "tesseract-multipass",
            "psm": 6,
            "data": {},
        }
        with patch.object(main, "_recognize_image", return_value=recognition):
            records = asyncio.run(plugin.capture_event(event))

        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["kind"], "image")
        self.assertEqual(records[0]["ocr_confidence"], 0.2)
        self.assertEqual(len(event.sent), 2)
        notice = next(item for item in event.sent if "LLM 审核暂不可用" in item)
        self.assertIn("重新发送更清晰或可解析的原文件", notice)
        self.assertNotIn("/审核文件", notice)

    def test_captures_and_extracts_docx_and_xlsx(self):
        try:
            from docx import Document
            from openpyxl import Workbook
        except ImportError as exc:
            self.skipTest(str(exc))

        docx_path = self.temp_path / "students.docx"
        document = Document()
        document.add_paragraph("姓名：周九")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "班级"
        table.cell(0, 1).text = "软件工程2301"
        document.save(docx_path)

        xlsx_path = self.temp_path / "students.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "名单"
        sheet.append(["姓名", "专业"])
        sheet.append(["吴十", "计算机科学"])
        workbook.save(xlsx_path)

        plugin = make_plugin(self.temp_path)
        event = FakeEvent([
            {"type": "file", "name": docx_path.name, "path": str(docx_path)},
            {"type": "file", "name": xlsx_path.name, "path": str(xlsx_path)},
        ])
        records = asyncio.run(plugin.capture_event(event))

        by_name = {record["name"]: record for record in records}
        self.assertEqual(by_name[docx_path.name]["kind"], "document")
        self.assertEqual(by_name[xlsx_path.name]["kind"], "spreadsheet")
        self.assertIn("周九", by_name[docx_path.name]["ocr_text"])
        self.assertIn("软件工程2301", by_name[docx_path.name]["ocr_text"])
        self.assertIn("吴十", by_name[xlsx_path.name]["ocr_text"])
        self.assertIn("计算机科学", by_name[xlsx_path.name]["ocr_text"])
        self.assertTrue(all(record["extraction_status"] == "extracted" for record in records))

    def test_gb18030_text_and_csv_are_decoded_and_reported(self):
        text_path = self.temp_path / "legacy-list.txt"
        csv_path = self.temp_path / "legacy-list.csv"
        text_path.write_bytes("姓名：张三\n班级：软件工程2401".encode("gb18030"))
        csv_path.write_bytes("姓名,专业\n李四,人工智能".encode("gb18030"))
        plugin = make_plugin(self.temp_path, capture_progress_notify=False)
        event = FakeEvent([
            {"type": "file", "name": text_path.name, "path": str(text_path)},
            {"type": "file", "name": csv_path.name, "path": str(csv_path)},
        ])

        records = asyncio.run(plugin.capture_event(event))

        by_name = {record["name"]: record for record in records}
        self.assertIn("张三", by_name[text_path.name]["ocr_text"])
        self.assertIn("李四", by_name[csv_path.name]["ocr_text"])
        self.assertTrue(all(record["encoding_issue"] for record in records))
        self.assertTrue(all(record["source_encoding"] == "gb18030" for record in records))
        self.assertEqual(len(event.sent), 1)
        self.assertIn(text_path.name, event.sent[0])
        self.assertIn(csv_path.name, event.sent[0])

    def test_utf8_mojibake_is_repaired_and_marked(self):
        path = self.temp_path / "mojibake.txt"
        mojibake = "姓名：张三".encode("utf-8").decode("latin-1")
        path.write_text(mojibake, encoding="utf-8")
        plugin = make_plugin(self.temp_path, capture_progress_notify=False)

        records = asyncio.run(plugin.capture_event(FakeEvent([
            {"type": "file", "name": path.name, "path": str(path)}
        ])))

        self.assertIn("姓名：张三", records[0]["ocr_text"])
        self.assertTrue(records[0]["encoding_issue"])
        self.assertEqual(records[0]["source_encoding"], "utf-8")
        self.assertEqual(records[0]["encoding_repair"], "latin-1->utf-8")

    def test_encoding_issue_notice_can_be_disabled(self):
        path = self.temp_path / "silent-gbk.txt"
        path.write_bytes("姓名：王五".encode("gb18030"))
        plugin = make_plugin(
            self.temp_path,
            capture_progress_notify=False,
            notify_encoding_issues=False,
        )
        event = FakeEvent([{"type": "file", "name": path.name, "path": str(path)}])

        records = asyncio.run(plugin.capture_event(event))

        self.assertTrue(records[0]["encoding_issue"])
        self.assertEqual(event.sent, [])

    def test_xlsx_extracts_all_sheets_and_formula_fallback(self):
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            self.skipTest(str(exc))

        path = self.temp_path / "multi-sheet.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.title = "一班"
        first.append(["姓名", "班级", "计算值"])
        first.append(["赵六", "软件工程2401", "=1+1"])
        second = workbook.create_sheet("二班")
        second.append(["姓名", "专业"])
        second.append(["孙七", "数据科学"])
        workbook.save(path)
        plugin = make_plugin(self.temp_path, capture_progress_notify=False)

        records = asyncio.run(plugin.capture_event(FakeEvent([
            {"type": "file", "name": path.name, "path": str(path)}
        ])))

        record = records[0]
        self.assertEqual(record["sheet_count"], 2)
        self.assertEqual(record["sheet_names"], ["一班", "二班"])
        self.assertIn("[工作表 1/2：一班]", record["ocr_text"])
        self.assertIn("[工作表 2/2：二班]", record["ocr_text"])
        self.assertIn("赵六", record["ocr_text"])
        self.assertIn("孙七", record["ocr_text"])
        self.assertIn("=1+1", record["ocr_text"])

    def test_docx_table_extraction_preserves_empty_and_merged_cells(self):
        try:
            from docx import Document
        except ImportError as exc:
            self.skipTest(str(exc))

        path = self.temp_path / "complex-table.docx"
        document = Document()
        table = document.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "班级"
        table.cell(0, 2).text = "备注"
        table.cell(1, 0).text = "张三"
        table.cell(1, 1).text = "软件工程2401"
        table.cell(1, 1).merge(table.cell(1, 2))
        document.save(path)

        text = main._extract_docx_text(path)

        self.assertIn("[表格 1]", text)
        self.assertIn("姓名 | 班级 | 备注", text)
        self.assertIn("张三", text)
        self.assertIn("软件工程2401", text)
        self.assertIn("[跨 2 列]", text)

    def test_all_files_use_multiple_vision_models_and_consensus(self):
        source = self.temp_path / "roster.txt"
        source.write_text("本地：张山，软件工程2401", encoding="utf-8")
        plugin = make_plugin(
            self.temp_path,
            vision_enabled=True,
            vision_model_targets=["provider-a::vision-a", "provider-b::vision-b"],
            vision_consensus_enabled=True,
            capture_progress_notify=False,
        )
        calls = []

        async def vision_call(_context, _path, _mime, _prompt, **kwargs):
            calls.append((kwargs.get("provider_name"), kwargs.get("model_name")))
            return f"{kwargs.get('model_name')}：姓名 张三 班级 软件工程2401", "mock-vision"

        async def consensus_call(_context, _prompt, **_kwargs):
            return "姓名 张三 | 班级 软件工程2401", "mock-consensus"

        with patch.object(main, "_prepare_vision_review_images", return_value=[source]), \
             patch.object(main, "_call_vision_model", side_effect=vision_call), \
             patch.object(main, "_call_text_model", side_effect=consensus_call):
            records = asyncio.run(plugin.capture_event(FakeEvent([
                {"type": "file", "name": source.name, "path": str(source)}
            ])))

        self.assertEqual(calls, [("provider-a", "vision-a"), ("provider-b", "vision-b")])
        self.assertEqual(records[0]["vision_models_completed"], 2)
        self.assertEqual(records[0]["vision_consensus_status"], "completed")
        self.assertEqual(records[0]["ocr_engine"], "vision-consensus")
        self.assertIn("张三", records[0]["ocr_text"])

    def test_capture_window_collects_multiple_events_before_processing(self):
        first = self.temp_path / "window-one.txt"
        second = self.temp_path / "window-two.txt"
        first.write_text("张三", encoding="utf-8")
        second.write_text("李四", encoding="utf-8")
        plugin = make_plugin(
            self.temp_path,
            vision_enabled=False,
            capture_window_enabled=True,
            capture_interval_seconds=1,
            capture_progress_notify=False,
        )

        async def scenario():
            await plugin.on_message(FakeEvent([{"type": "file", "name": first.name, "path": str(first)}]))
            await plugin.on_message(FakeEvent([{"type": "file", "name": second.name, "path": str(second)}]))
            collecting = plugin.latest_progress()
            task = plugin._capture_window_task
            self.assertIsNotNone(task)
            await task
            return collecting, plugin.latest_progress()

        collecting, completed = asyncio.run(scenario())

        self.assertEqual(collecting["status"], "collecting")
        self.assertEqual(collecting["total"], 2)
        self.assertEqual(completed["status"], "completed")
        self.assertEqual(completed["completed"], 2)
        self.assertEqual({item["name"] for item in plugin.store.list()}, {first.name, second.name})

    def test_web_panel_lists_and_deletes_captured_files(self):
        plugin = make_plugin(self.temp_path, panel_access_token="secret-token")
        first = plugin.store.add_bytes("one.docx", b"one", kind="document", mime="application/octet-stream", source={})
        second = plugin.store.add_bytes("two.xlsx", b"two", kind="spreadsheet", mime="application/octet-stream", source={})
        app = main._build_web_app(plugin)
        if app is None:
            self.skipTest("starlette is not installed")

        async def request(method, path, payload=None):
            body = json.dumps(payload or {}).encode("utf-8") if payload is not None else b""
            sent = []
            delivered = False

            async def receive():
                nonlocal delivered
                if delivered:
                    return {"type": "http.disconnect"}
                delivered = True
                return {"type": "http.request", "body": body, "more_body": False}

            async def send(message):
                sent.append(message)

            scope = {
                "type": "http",
                "asgi": {"version": "3.0"},
                "http_version": "1.1",
                "method": method,
                "scheme": "http",
                "path": path,
                "raw_path": path.encode("ascii"),
                "query_string": b"",
                "root_path": "",
                "headers": [(b"content-type", b"application/json"), (b"authorization", b"Bearer secret-token")],
                "client": ("127.0.0.1", 12345),
                "server": ("test", 80),
            }
            await app(scope, receive, send)
            status = next(message["status"] for message in sent if message["type"] == "http.response.start")
            response_body = b"".join(message.get("body", b"") for message in sent if message["type"] == "http.response.body")
            return status, json.loads(response_body or b"{}")

        status, listing = asyncio.run(request("GET", "/files"))
        self.assertEqual(status, 200)
        self.assertEqual({item["id"] for item in listing["files"]}, {first["id"], second["id"]})
        self.assertTrue(all("path" not in item for item in listing["files"]))

        progress_id = plugin._start_progress(FakeEvent([]), [])
        plugin._update_progress(progress_id, advance=False, status="completed")
        status, progress = asyncio.run(request("GET", "/progress"))
        self.assertEqual(status, 200)
        self.assertEqual(progress["progress"]["id"], progress_id)
        self.assertEqual(progress["progress"]["percent"], 100)

        status, result = asyncio.run(request("POST", "/files/bulk-delete", {"ids": [first["id"]]}))
        self.assertEqual(status, 200)
        self.assertEqual(result["deleted"], [first["id"]])
        self.assertIsNone(plugin.store.get(first["id"]))
        self.assertIsNotNone(plugin.store.get(second["id"]))

    def test_web_panel_mounts_on_asgi_host_and_command_returns_path(self):
        class HostApp:
            def __init__(self):
                self.mounts = []

            def mount(self, path, app, name=None):
                self.mounts.append((path, app, name))

        class HostContext:
            def __init__(self):
                self.app = HostApp()

        context = HostContext()
        plugin = main.NameSearcherPlugin(
            context,
            {"storage_dir": str(self.temp_path / "files"), "panel_mount_path": "/api/name-searcher", "restrict_chat_scope": False},
        )

        self.assertEqual(len(context.app.mounts), 1)
        self.assertEqual(context.app.mounts[0][0], "/api/name-searcher")
        event = FakeEvent([], text="/文件面板")
        output = asyncio.run(collect_async_generator(plugin.file_panel(event)))
        self.assertEqual(len(output), 1)
        self.assertIn("文件管理面板已注册到 AstrBot 网页端", output[0])
        self.assertIn("插件 -> 姓名文件查找器 -> 文件管理与预览", output[0])
        self.assertIn("旧版兼容地址：/api/name-searcher/", output[0])

    def test_official_plugin_page_apis_are_registered(self):
        class PageContext:
            def __init__(self):
                self.routes = []

            def register_web_api(self, route, handler, methods, description):
                self.routes.append((route, handler, methods, description))

        context = PageContext()
        plugin = main.NameSearcherPlugin(
            context,
            {"storage_dir": str(self.temp_path / "files")},
        )
        routes = {(route, tuple(methods)) for route, _handler, methods, _description in context.routes}

        self.assertIn((f"/{main.PLUGIN_ID}/files", ("GET",)), routes)
        self.assertIn((f"/{main.PLUGIN_ID}/files/delete", ("POST",)), routes)
        self.assertIn((f"/{main.PLUGIN_ID}/files/upload", ("POST",)), routes)
        self.assertIn((f"/{main.PLUGIN_ID}/files/<artifact_id>/content", ("GET",)), routes)
        self.assertTrue(main.PLUGIN_PAGE_ROOT.joinpath("index.html").exists())
        response = asyncio.run(plugin._page_files())
        self.assertEqual(response, {"files": [], "count": 0})

    def test_bridge_image_preview_returns_data_uri_without_local_path(self):
        plugin = make_plugin(self.temp_path)
        record = plugin.store.add_bytes(
            "preview.png",
            b"fake-png-payload",
            kind="image",
            mime="image/png",
            source={},
        )

        payload, status = main._preview_payload(plugin.store, record["id"])

        self.assertEqual(status, 200)
        self.assertTrue(payload["data_url"].startswith("data:image/png;base64,"))
        self.assertNotIn("path", payload)

    def test_store_bulk_delete_and_clear_remove_downloaded_files(self):
        plugin = make_plugin(self.temp_path)
        records = [
            plugin.store.add_bytes(
                f"file-{index}.txt",
                f"content-{index}".encode("ascii"),
                kind="document",
                mime="text/plain",
                source={},
            )
            for index in range(3)
        ]
        paths = [plugin.store.path_for(record) for record in records]
        self.assertTrue(all(path is not None and path.exists() for path in paths))

        deleted = plugin.store.delete([records[0]["id"], records[2]["id"]])

        self.assertCountEqual(deleted, [records[0]["id"], records[2]["id"]])
        self.assertFalse(paths[0].exists())
        self.assertTrue(paths[1].exists())
        self.assertFalse(paths[2].exists())
        self.assertEqual([item["id"] for item in plugin.store.list()], [records[1]["id"]])

        self.assertEqual(plugin.store.clear(), 1)
        self.assertFalse(paths[1].exists())
        self.assertEqual(plugin.store.list(), [])
        self.assertEqual(json.loads(plugin.store.index_path.read_text(encoding="utf-8")), [])


if __name__ == "__main__":
    unittest.main()
